# =============================================================================
# region SCRIPT DETAILS
# =============================================================================
# CSP CST 627 100  Deep Learning Neural Networks
# Week 05 Assignment: Latency, Memory, Accuracy: A Measured Trade-off Board
# Objective: Build the benchmark workbook (Excel) from the measured results.
#
# Usage:
#   python scripts/build_workbook.py                 # every full run in results/ (VM CPU and/or T4)
#   python scripts/build_workbook.py --tags cpu      # only the listed tags
# The assignment script calls build() as its last compute step.
# =============================================================================
# endregion
# =============================================================================
# Step 1: Load the Tagged Results
# Reads results/results_<tag>.json, manifest_<tag>.json and the raw run CSVs for each tag.
# Quick smoke-test runs (tag ending _quick) are only included when named explicitly.
# =============================================================================
"""
https://docs.python.org/3/library/json.html
https://pandas.pydata.org/docs/reference/api/pandas.read_csv.html
"""

import argparse
import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col

ROOT = Path(__file__).resolve().parents[1]
TAG_ORDER = ["cpu", "t4gpu", "gpu"]
TAG_LABEL = {"cpu": "VM / CPU only", "t4gpu": "Colab / T4 GPU", "gpu": "Other GPU"}
HF = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")


def discover_tags(root, explicit=None):
    if explicit:
        return explicit
    found = [p.stem.replace("results_", "") for p in (root / "results").glob("results_*.json")]
    return sorted((t for t in found if not t.endswith("_quick")), key=lambda t: TAG_ORDER.index(t) if t in TAG_ORDER else 99)


def load(root, tags):
    runs = {}
    for t in tags:
        r = root / "results"
        runs[t] = dict(
            results=json.loads((r / f"results_{t}.json").read_text(encoding="utf-8")),
            manifest=json.loads((r / f"manifest_{t}.json").read_text(encoding="utf-8")),
            acc=pd.read_csv(r / f"accuracy_runs_{t}.csv"),
            lat=pd.read_csv(r / f"latency_runs_{t}.csv"))
    return runs


# =============================================================================
# Step 2: Build the Workbook
# Board rows are live `AVERAGEIFS` formulas over the raw Acc_Runs / Latency_Runs sheets (two
# criteria: tag and config), so every board number can be traced to individual runs. Standard
# deviations are written as values because Excel has no STDEVIFS. The ScalingLaw sheet's
# assumptions (peak TFLOPs, MFU, tokens per parameter, scale-up factor) are yellow input cells.
# =============================================================================
"""
https://openpyxl.readthedocs.io/en/stable/
https://support.microsoft.com/en-us/office/averageifs-function-48910c45-1fc0-4389-a028-f7c5c3001690
"""


def header(ws, row, cols, width=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row, j, c)
        cell.font, cell.fill = HF, HFILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[col(j)].width = width or max(12, min(34, len(str(c)) + 4))


def _nan(x):
    return None if x is None or (isinstance(x, float) and math.isnan(x)) else x


def sheet_runs(wb, runs):
    wa = wb.create_sheet("Acc_Runs")
    header(wa, 1, ["tag", "config", "seed", "accuracy", "n_test", "ci95"])
    wl = wb.create_sheet("Latency_Runs")
    header(wl, 1, ["tag", "config", "device", "trial", "p50_ms", "p95_ms", "mean_ms", "std_ms", "n_iter"])
    for t, d in runs.items():
        for r in d["acc"][["tag", "config", "seed", "accuracy", "n_test", "ci95"]].itertuples(index=False):
            wa.append(list(r))
        for r in d["lat"][["tag", "config", "device", "trial", "p50_ms", "p95_ms", "mean_ms", "std_ms", "n_iter"]].itertuples(index=False):
            wl.append(list(r))
    return wa.max_row, wl.max_row


def sheet_board(wb, runs, nA, nL, notes):
    ws = wb.create_sheet("Board", 0)
    ws["A1"] = "Brightline Logistics - Edge Model Trade-off Board"
    ws["A1"].font = Font(bold=True, size=14)
    desc = []
    for t, d in runs.items():
        h, c = d["manifest"]["hardware"], d["results"]["cfg"]
        desc.append(f"[{t}] {h.get('host_label')} | {h.get('gpu_name') or h.get('cpu_model')} | torch {h['torch']} | seeds {c['seeds']} | "
                    f"{c['trials']} trials x {c['n_iter']} timed forwards, batch 1, n_test={c['n_test']}")
    ws["A2"] = "   ||   ".join(desc) + ("   ** QUICK SMOKE RUN - NOT FINAL NUMBERS **" if any(d["results"]["quick"] for d in runs.values()) else "")
    cols = ["tag", "config", "technique", "device", "precision", "params", "size MB", "accuracy (mean of seeds)", "accuracy std (seeds)",
            "acc 95% CI +/- (binomial)", "p50 ms (mean of trials)", "p50 std across trials", "p95 ms (mean of trials)", "p95 std across trials",
            "peak GPU MB (above floor)", "CPU RSS delta MB", "reference config", "acc delta vs ref", "p50 speedup vs ref (x)", "size ratio vs ref", "note",
            "peak GPU MB (absolute)", "GPU floor MB (allocated before load)"]
    header(ws, 4, cols)
    rowof, r = {}, 5
    for t, d in runs.items():
        for key, c in d["results"]["configs"].items():
            rowof[(t, key)] = r
            r += 1
    for (t, key), r in rowof.items():
        c = runs[t]["results"]["configs"][key]
        ws.cell(r, 1, t); ws.cell(r, 2, key); ws.cell(r, 3, c["label"]); ws.cell(r, 4, c["device"]); ws.cell(r, 5, c["precision"])
        ws.cell(r, 6, c["params"]); ws.cell(r, 7, round(c["size_mb"], 4))
        ws.cell(r, 8, f"=AVERAGEIFS(Acc_Runs!D2:D{nA},Acc_Runs!A2:A{nA},A{r},Acc_Runs!B2:B{nA},B{r})")
        ws.cell(r, 9, c["acc_std"])
        ws.cell(r, 10, f"=AVERAGEIFS(Acc_Runs!F2:F{nA},Acc_Runs!A2:A{nA},A{r},Acc_Runs!B2:B{nA},B{r})")
        ws.cell(r, 11, f"=AVERAGEIFS(Latency_Runs!E2:E{nL},Latency_Runs!A2:A{nL},A{r},Latency_Runs!B2:B{nL},B{r})")
        ws.cell(r, 12, c["p50_std"])
        ws.cell(r, 13, f"=AVERAGEIFS(Latency_Runs!F2:F{nL},Latency_Runs!A2:A{nL},A{r},Latency_Runs!B2:B{nL},B{r})")
        ws.cell(r, 14, c["p95_std"])
        ws.cell(r, 15, _nan(None if c["peak_gpu_mb"] is None else round(c["peak_gpu_mb"], 3)))
        ws.cell(r, 16, round(c["rss_delta_mb"], 2))
        ws.cell(r, 17, c["ref"] or "-")
        if c["ref"] and (t, c["ref"]) in rowof:
            rr = rowof[(t, c["ref"])]
            ws.cell(r, 18, f"=H{r}-H{rr}"); ws.cell(r, 19, f"=K{rr}/K{r}"); ws.cell(r, 20, f"=G{r}/G{rr}")
        ws.cell(r, 21, notes.get(key, ""))
        ws.cell(r, 22, _nan(None if c.get("peak_gpu_abs_mb") is None else round(c["peak_gpu_abs_mb"], 3)))
        ws.cell(r, 23, _nan(None if c.get("gpu_floor_mb") is None else round(c["gpu_floor_mb"], 3)))
        for cc in (8, 9, 10, 18):
            ws.cell(r, cc).number_format = "0.0000"
        for cc in (11, 12, 13, 14, 19, 20):
            ws.cell(r, cc).number_format = "0.000"
    ws.freeze_panes = "C5"
    ws.column_dimensions["C"].width = 46
    ws.column_dimensions["U"].width = 60
    return ws, rowof


def sheet_cross(wb, runs, rowof):
    """Same config on both machines: the Intel-VM-vs-T4 comparison."""
    if "cpu" not in runs or not any(t != "cpu" for t in runs):
        return
    gt = next(t for t in runs if t != "cpu")
    ws = wb.create_sheet("VM_vs_GPU", 1)
    ws["A1"] = f"Same code, same seeds, same budgets: Intel VM (tag cpu) vs {TAG_LABEL.get(gt, gt)} (tag {gt})"
    ws["A1"].font = Font(bold=True, size=13)
    header(ws, 3, ["config", "VM acc", "GPU-run acc", "acc diff", "VM p50 ms", "GPU-run p50 ms", "p50 speedup (VM / GPU run)",
                   "VM p95 ms", "GPU-run p95 ms", "p95 speedup", "where the GPU run executed it"])
    i = 4
    for key in runs["cpu"]["results"]["configs"]:
        if (gt, key) not in rowof:
            continue
        a, b = rowof[("cpu", key)], rowof[(gt, key)]
        ws.cell(i, 1, key); ws.cell(i, 2, f"=Board!H{a}"); ws.cell(i, 3, f"=Board!H{b}"); ws.cell(i, 4, f"=C{i}-B{i}")
        ws.cell(i, 5, f"=Board!K{a}"); ws.cell(i, 6, f"=Board!K{b}"); ws.cell(i, 7, f"=E{i}/F{i}")
        ws.cell(i, 8, f"=Board!M{a}"); ws.cell(i, 9, f"=Board!M{b}"); ws.cell(i, 10, f"=H{i}/I{i}")
        ws.cell(i, 11, runs[gt]["results"]["configs"][key]["device"])
        for cc in (2, 3, 4):
            ws.cell(i, cc).number_format = "0.0000"
        for cc in range(5, 11):
            ws.cell(i, cc).number_format = "0.000"
        i += 1
    ws.cell(i + 1, 1, "Quality should agree across hardware (small drift = floating-point order). Latency is what the hardware changes. "
                      "Note the INT8 and CPU rows run on the CPU of whichever machine this was: the T4 run's CPU is Colab's 2-vCPU Xeon, not the T4.")
    ws.column_dimensions["A"].width = 24


def sheet_scaling(wb, runs):
    t0 = next(iter(runs))
    scal = pd.DataFrame(runs[t0]["results"]["scaling"])
    ws = wb.create_sheet("ScalingLaw")
    ws["A1"] = "Scaling-law compute estimate (6ND and Chinchilla-style token budget)"
    ws["A1"].font = Font(bold=True, size=13)
    for i, (k, v) in enumerate([("GPU peak FP16 TFLOPs (Colab T4, vendor spec)", 65), ("Assumed model-FLOPs utilisation (MFU)", 0.30),
                                ("Scale-up factor for params", 10), ("Chinchilla tokens per param (D* = k*N)", 20),
                                ("Colab free-tier session cap (hours, approximate)", 12)]):
        ws.cell(3 + i, 1, k)
        ws.cell(3 + i, 2, v).fill = INPUT_FILL
    TF, MFU, X, K, CAP = "$B$3", "$B$4", "$B$5", "$B$6", "$B$7"
    header(ws, 10, ["model", "params N", "unique train samples", "tokens/sample (24 text + 64 patches)", "train samples seen (epochs x N)",
                    "actual tokens D = samples seen x tokens/sample", "actual train FLOPs 6ND", "measured fwd FLOPs/sample",
                    "cross-check: 3 x fwd x samples seen", "6ND / measured ratio", "Chinchilla D* = k*N", "D / D*", "unique tokens / D*",
                    "10x params N'", "10x Chinchilla tokens D'=k*N'", "10x train FLOPs 6N'D'", "10x GPU-hours on T4", "10x fits in session cap?",
                    "years on one T4 (reference rows)"])
    for i, s in enumerate(scal.itertuples(index=False)):
        r = 11 + i
        ws.cell(r, 1, s.model); ws.cell(r, 2, s.params); ws.cell(r, 3, s.unique_train_samples); ws.cell(r, 4, s.tokens_per_sample)
        ws.cell(r, 5, s.train_samples_seen); ws.cell(r, 6, f"=E{r}*D{r}"); ws.cell(r, 7, f"=6*B{r}*F{r}"); ws.cell(r, 8, s.fwd_flops_per_sample)
        ws.cell(r, 9, f"=3*H{r}*E{r}"); ws.cell(r, 10, f"=G{r}/I{r}"); ws.cell(r, 11, f"={K}*B{r}"); ws.cell(r, 12, f"=F{r}/K{r}")
        ws.cell(r, 13, f"=C{r}*D{r}/K{r}"); ws.cell(r, 14, f"={X}*B{r}"); ws.cell(r, 15, f"={K}*N{r}"); ws.cell(r, 16, f"=6*N{r}*O{r}")
        ws.cell(r, 17, f"=P{r}/({TF}*10^12*{MFU})/3600"); ws.cell(r, 18, f'=IF(Q{r}<={CAP},"yes","NO")')
    rr = 11 + len(scal) + 1
    ws.cell(rr, 1, "Reference rows: what a single free GPU cannot answer (illustrative, not measured):").font = Font(bold=True)
    for j, (nm, n) in enumerate([("Reference: 0.9B-parameter audio model (MOSS-Transcribe-Diarize class)", 0.9e9), ("Reference: 10x of that (9B)", 9e9)]):
        r = rr + 1 + j
        ws.cell(r, 1, nm); ws.cell(r, 2, n); ws.cell(r, 11, f"={K}*B{r}"); ws.cell(r, 7, f"=6*B{r}*K{r}")
        ws.cell(r, 17, f"=G{r}/({TF}*10^12*{MFU})/3600"); ws.cell(r, 18, f'=IF(Q{r}<={CAP},"yes","NO")'); ws.cell(r, 19, f"=Q{r}/24/365")
    notes = ["How to read this:",
             "- 6ND: training FLOPs ~ 6 x params x tokens (2ND forward + 4ND backward). Approximate here: the CNN and a 24-token transformer are not transformer-dominated, hence the cross-check against measured FLOPs (column J).",
             "- Chinchilla heuristic: compute-optimal tokens D* ~ 20 x N. Our actual D counts repeated epochs; 'unique tokens / D*' shows how data-limited the run is.",
             "- The 10x rows apply the same heuristic. For these tiny models compute is trivial, so the GPU budget is NOT the binding constraint; the realism of the DATA is.",
             "- The reference rows bound the single-GPU question: a 0.9B model at Chinchilla tokens needs on the order of two months of T4 time (about 1,400 GPU-hours, against a ~12 h session), and its 10x version about 16 years. We can fine-tune or quantize such a model, but not answer how it would scale or how a 10x version would train.",
             "- Peak TFLOPs and MFU are assumptions (yellow cells). Change them and every dependent cell updates."]
    for i, t in enumerate(notes):
        ws.cell(rr + 5 + i, 1, t)
    ws.column_dimensions["A"].width = 62
    for r in range(11, rr + 4):
        for c in (2, 6, 7, 8, 9, 11, 14, 15, 16, 17):
            ws.cell(r, c).number_format = "0.00E+00"


def sheet_hardware(wb, runs):
    ws = wb.create_sheet("Hardware")
    keys = ["host_class", "host_label", "platform", "python", "torch", "torch_cuda_build", "device", "gpu_name", "vram_gb", "compute_capability",
            "driver", "cpu_model", "cpu_count", "ram_gib", "virtualized", "vm_product", "numpy", "pandas", "matplotlib"]
    header(ws, 1, ["field"] + list(runs), 30)
    for i, k in enumerate(keys, 2):
        ws.cell(i, 1, k)
        for j, t in enumerate(runs, 2):
            ws.cell(i, j, str(runs[t]["manifest"]["hardware"].get(k)))
    ws.cell(len(keys) + 3, 1, "total runtime (s)")
    for j, t in enumerate(runs, 2):
        ws.cell(len(keys) + 3, j, runs[t]["manifest"]["total_runtime_s"])


def sheet_method(wb, runs):
    ws = wb.create_sheet("Methodology", 1)
    ws.column_dimensions["A"].width = 160
    c = next(iter(runs.values()))["results"]["cfg"]
    lines = [
        "METHODOLOGY AND LIMITS", "",
        "Task: 6-class claim-queue routing from (bill-of-lading image, driver-note transcript). Label = cargo condition (intact/crushed/wet, from the image) x receiver outcome (accepted/refused, from the note).",
        f"Data: synthetic, generated by the assignment script with a fixed seed. Image: noisy page, low-contrast cue (depth {c['cue_depth']}), faint decoy scratches and smudges. Note: 60% carry a decoy clause about a different load with the opposite outcome. "
        f"{int(c['cue_noise_img'] * 100)}% of image cues and {int(c['cue_noise_note'] * 100)}% of note cues are wrong, so the ceiling is roughly 0.79. Same data for every configuration. The test set is never used for training or selection.",
        "Driver notes are text transcripts. No audio model is in the loop, so speech-recognition error is NOT measured (a stated gap).",
        "Accuracy: mean over training seeds on a fixed test set; seed std, plus a binomial 95% interval that covers test-set sampling error only.",
        "Latency: batch size 1; warm-up before every timed trial; torch.cuda.synchronize() before and after each timed forward; repeated trials; p50/p95 per trial; variance reported across trials.",
        "Confounds controlled: input transfer excluded from the timer; CPU thread count pinned; configs visited in shuffled order within each trial round; each config's memory measured alone on a fresh load.",
        "Memory: GPU peak = torch.cuda.max_memory_allocated (excludes the CUDA context). CPU: serialized model size plus a coarse process-RSS delta.",
        "Techniques: dynamic INT8 quantization of Linear layers (CPU only); FP16 on GPU; knowledge distillation with a labels-only control; linear-attention variant; late vs early fusion (early fusion also trained for 2x epochs to separate 'worse' from 'slower to train').",
        "Hardware: runs are tagged by machine. 'cpu' = the Intel VM (CPU only); 't4gpu' = Colab T4. GPU-labelled configs run on the primary device, so on the VM they run on the CPU.", "",
        "KNOWN LIMITS (do not over-claim):",
        "- Synthetic data: absolute accuracy says nothing about real scanned bills of lading; only the relative ordering under controlled conditions.",
        "- Neither machine is the in-cab tablet. CPU INT8 on a VM or on 2 Colab vCPUs is a proxy; real Android NPU/XNNPACK latencies will differ.",
        "- A virtual machine adds scheduling jitter, which is why p95 and trial-to-trial spread are reported.",
        "- GPU memory is the peak above the allocation present just before each config loads (weights, activations, and any workspace that config allocates). Library workspaces such as cuBLAS stay allocated after training and would add the same offset to every row; the absolute peak and the floor are in columns V and W.",
        "- Sequence length 24 is too short for linear attention to pay off; the variant tests that, not long-context behaviour.",
        "- Dynamic INT8 leaves convolutions in FP32; full static INT8 or QAT could behave differently.",
        "- The free Colab GPU is shared; absolute latencies vary between sessions, so compare configs within one run, not across runs.",
        f"- {len(c['seeds'])} training seeds: accuracy gaps smaller than about one seed std are not distinguishable.", "",
        "Reproduce: run the notebook on a Colab T4 (COLAB.md) and the installer script on the VM; both write results/*_<tag>.* and rebuild this workbook."]
    for i, t in enumerate(lines, 1):
        ws.cell(i, 1, t).alignment = Alignment(wrap_text=True)
    ws["A1"].font = Font(bold=True, size=13)


NOTES = {
    "baseline_fp32": "Reference for the primary device", "baseline_fp16": "Same weights as baseline",
    "linear_attn": "Sequence length is only 24: no asymptotic advantage expected", "early_fusion": "Image patches + words in one transformer",
    "early_fusion_2x": "Same model, double the epochs: separates 'worse' from 'slower to train'", "student_kd": "Width 16, 1 layer, distilled from the baseline",
    "student_nokd": "Same student from labels only: isolates the effect of KD", "cpu_baseline_fp32": "CPU reference when the primary device is a GPU",
    "cpu_baseline_int8": "Dynamic INT8 on Linear layers; convolutions stay FP32"}


def chart(wb, runs, board_ws, out_dir):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        fig, axes = plt.subplots(1, len(runs), figsize=(6.2 * len(runs), 4.6), squeeze=False)
        for ax, (t, d) in zip(axes[0], runs.items()):
            for k, c in d["results"]["configs"].items():
                ax.scatter(c["p50_mean"], c["acc_mean"], s=40 + 4 * c["params"] / 1000, alpha=.7, marker="s" if c["device"] == "cpu" else "o")
                ax.annotate(k.replace("_", " "), (c["p50_mean"], c["acc_mean"]), fontsize=7)
            ax.set_xlabel("p50 latency, batch 1 (ms)")
            ax.set_ylabel("test accuracy (mean of seeds)")
            ax.set_title(f"{d['manifest']['hardware'].get('host_label')}  [{t}]", fontsize=9)
        fig.tight_layout()
        path = out_dir / "board_tradeoff.png"
        fig.savefig(path, dpi=130)
        plt.close(fig)
        board_ws.add_image(XLImage(str(path)), f"B{board_ws.max_row + 3}")
    except Exception as exc:  # the chart is a convenience; never block the workbook on it
        print("[workbook] chart skipped:", exc)


def build(root=ROOT, tags=None, out=None):
    root = Path(root)
    tags = discover_tags(root, tags)
    if not tags:
        raise SystemExit("No results found in results/. Run the assignment script first.")
    runs = load(root, tags)
    wb = Workbook()
    wb.remove(wb.active)
    nA, nL = sheet_runs(wb, runs)
    board, rowof = sheet_board(wb, runs, nA, nL, NOTES)
    sheet_cross(wb, runs, rowof)
    sheet_method(wb, runs)
    sheet_scaling(wb, runs)
    sheet_hardware(wb, runs)
    figdir = root / "results" / "figs_workbook"
    figdir.mkdir(exist_ok=True)
    chart(wb, runs, board, figdir)
    out = Path(out or root / "results" / "Week05_Benchmark_Workbook.xlsx")
    wb.save(out)
    print(f"[workbook] saved {out.relative_to(root) if out.is_relative_to(root) else out}  (tags: {', '.join(tags)})")
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Build the benchmark workbook from results/*.json")
    ap.add_argument("--tags", nargs="*", default=None, help="tags to include (default: every full run)")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    build(ROOT, a.tags, a.out)
